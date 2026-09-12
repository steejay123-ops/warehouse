from django.test import TestCase
from accounts.serializers import UserSerializer
from accounts.models import CustomUser

class UserSerializerValidationTests(TestCase):
    def test_create_user_without_email_succeeds(self):
        data = {
            'username': 'user_no_email',
            'first_name': 'علی',
            'last_name': 'رضایی',
            'phone_number': '09123456789',
        }
        serializer = UserSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        user = serializer.save()
        self.assertEqual(user.email, '')

    def test_create_user_with_null_email_succeeds(self):
        data = {
            'username': 'user_null_email',
            'first_name': 'سارا',
            'last_name': 'محمدی',
            'phone_number': '09123456789',
            'email': None
        }
        serializer = UserSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        user = serializer.save()
        self.assertEqual(user.email, '')

    def test_create_user_without_phone_fails(self):
        data = {
            'username': 'user_no_phone',
            'first_name': 'مهدی',
            'last_name': 'کریمی',
        }
        serializer = UserSerializer(data=data)
        self.assertFalse(serializer.is_valid())
        self.assertIn('phone_number', serializer.errors)

    def test_create_user_with_invalid_phone_fails(self):
        data = {
            'username': 'user_bad_phone',
            'first_name': 'مهدی',
            'last_name': 'کریمی',
            'phone_number': '02188776655', # Not mobile
        }
        serializer = UserSerializer(data=data)
        self.assertFalse(serializer.is_valid())
        self.assertIn('phone_number', serializer.errors)

    def test_create_user_with_iranian_mobile_normalization_succeeds(self):
        data = {
            'username': 'user_iran_phone',
            'first_name': 'مهدی',
            'last_name': 'کریمی',
            'phone_number': '+989123456789',
        }
        serializer = UserSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        user = serializer.save()
        self.assertEqual(user.phone_number, '09123456789')

    def test_create_user_with_invalid_email_fails(self):
        data = {
            'username': 'bad_email_user',
            'first_name': 'رضا',
            'last_name': 'محمدی',
            'phone_number': '09121112233',
            'email': 'not-a-valid-email',
        }
        serializer = UserSerializer(data=data)
        self.assertFalse(serializer.is_valid())
        self.assertIn('email', serializer.errors)
        self.assertIn('نامعتبر', str(serializer.errors['email']))

    def test_create_user_with_valid_email_succeeds(self):
        data = {
            'username': 'good_email_user',
            'first_name': 'رضا',
            'last_name': 'محمدی',
            'phone_number': '09121112233',
            'email': 'reza.mohammadi@example.com',
        }
        serializer = UserSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)
        self.assertEqual(serializer.validated_data['email'], 'reza.mohammadi@example.com')

    def test_create_user_with_short_password_fails(self):
        data = {
            'username': 'short_pwd_user',
            'first_name': 'رضا',
            'last_name': 'محمدی',
            'phone_number': '09121112233',
            'password': '123',
        }
        serializer = UserSerializer(data=data)
        self.assertFalse(serializer.is_valid())
        self.assertIn('password', serializer.errors)
        self.assertIn('۶ کاراکتر', str(serializer.errors['password']))

    def test_create_user_with_valid_password_succeeds(self):
        data = {
            'username': 'good_pwd_user',
            'first_name': 'رضا',
            'last_name': 'محمدی',
            'phone_number': '09121112233',
            'password': 'StrongPassword123!',
        }
        serializer = UserSerializer(data=data)
        self.assertTrue(serializer.is_valid(), serializer.errors)



import io
import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient
from django.contrib.auth.models import Permission
from accounts.models import CustomRole

def create_excel_upload_file(headers, key_row, rows, filename="test.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(key_row)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return SimpleUploadedFile(filename, buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class UserExcelImportTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = CustomUser.objects.create_superuser('admin_test_excel', 'admin@test.com', 'pass1234')
        self.client.force_authenticate(user=self.admin)

    def test_clear_empty_roles_and_warehouses_on_update(self):
        role = CustomRole.objects.create(name='accountant_test', title='حسابدار', color='#4f46e5')
        user = CustomUser.objects.create(
            username='user_sync_test',
            first_name='سارا',
            last_name='احمدی',
            phone_number='09121112233'
        )
        user.groups.add(role)
        self.assertEqual(user.groups.count(), 1)

        headers = ['نام', 'نام خانوادگی', 'شناسه ورود', 'کد ملی', 'تلفن', 'ایمیل', 'گروه خونی', 'تماس اضطراری', 'منطقه عملیاتی', 'شرکت متبوع', 'آدرس', 'نقش‌ها', 'انبارها', 'فعال']
        key_row = ['first_name', 'last_name', 'username', 'national_code', 'phone_number', 'email', 'blood_type', 'emergency_contact', 'operational_zone', 'company', 'address', 'roles', 'warehouses', 'is_active']
        rows = [
            ['سارا', 'احمدی', 'user_sync_test', '', '09121112233', '', '', '', '', '', '', '', '', 'بله']
        ]
        excel_file = create_excel_upload_file(headers, key_row, rows)
        res = self.client.post('/api/accounts/users/import_excel/', {'file': excel_file, 'update_existing': 'true'}, format='multipart')
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data['success'])
        self.assertEqual(res.data['summary']['updated'], 1)

        user.refresh_from_db()
        self.assertEqual(user.groups.count(), 0)

    def test_atomic_rollback_on_user_import_error(self):
        headers = ['نام', 'نام خانوادگی', 'شناسه ورود', 'کد ملی', 'تلفن', 'ایمیل', 'گروه خونی', 'تماس اضطراری', 'منطقه عملیاتی', 'شرکت متبوع', 'آدرس', 'نقش‌ها', 'انبارها', 'فعال']
        key_row = ['first_name', 'last_name', 'username', 'national_code', 'phone_number', 'email', 'blood_type', 'emergency_contact', 'operational_zone', 'company', 'address', 'roles', 'warehouses', 'is_active']
        rows = [
            ['کاربر', 'یک', 'rollback_user_1', '', '09121110001', '', '', '', '', '', '', '', '', 'بله'],
            ['کاربر', 'دو', 'rollback_user_2', '', '09121110002', '', '', '', '', '', '', '', '', 'بله'],
        ]
        excel_file = create_excel_upload_file(headers, key_row, rows)

        from unittest.mock import patch
        with patch.object(CustomUser, 'save', side_effect=[None, Exception("Simulated DB Crash")]):
            res = self.client.post('/api/accounts/users/import_excel/', {'file': excel_file}, format='multipart')
            self.assertEqual(res.status_code, 400)
            self.assertFalse(res.data['success'])
            self.assertIn('تراکنش کاملاً رول‌بک شد', res.data['errors'][0]['message'])

        self.assertFalse(CustomUser.objects.filter(username='rollback_user_1').exists())
        self.assertFalse(CustomUser.objects.filter(username='rollback_user_2').exists())


class RoleExcelImportTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = CustomUser.objects.create_superuser('admin_role_test', 'admin@role.com', 'pass1234')
        self.client.force_authenticate(user=self.admin)

    def test_circular_dependency_handled_gracefully_without_500(self):
        headers = ['نام یکتا (name)', 'عنوان فارسی (title)', 'رنگ سازمانی (color)', 'نقش والد (parent)', 'مجوزها (permissions)']
        key_row = ['name', 'title', 'color', 'parent', 'permissions']
        rows = [
            ['role_alpha', 'نقش آلفا', '#4f46e5', 'role_beta', ''],
            ['role_beta', 'نقش بتا', '#06b6d4', 'role_alpha', ''],
        ]
        excel_file = create_excel_upload_file(headers, key_row, rows)
        res = self.client.post('/api/accounts/roles/import_excel/', {'file': excel_file}, format='multipart')

        # MUST NOT be 500!
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data['success'])
        parent_errors = [e for e in res.data['errors'] if e.get('field') == 'parent']
        self.assertTrue(len(parent_errors) > 0)
        self.assertIn('چرخه‌ای', parent_errors[0]['message'])

    def test_clear_empty_permissions_and_parent_on_update(self):
        parent_role = CustomRole.objects.create(name='parent_lead', title='سرپرست کل', color='#10b981')
        perm = Permission.objects.filter(codename='view_sys_dashboard').first()
        child_role = CustomRole.objects.create(name='child_dev', title='توسعه دهنده', color='#6366f1', parent=parent_role)
        if perm:
            child_role.permissions.add(perm)

        self.assertIsNotNone(child_role.parent)
        if perm:
            self.assertEqual(child_role.permissions.count(), 1)

        headers = ['نام یکتا (name)', 'عنوان فارسی (title)', 'رنگ سازمانی (color)', 'نقش والد (parent)', 'مجوزها (permissions)']
        key_row = ['name', 'title', 'color', 'parent', 'permissions']
        rows = [
            ['child_dev', 'توسعه دهنده بروز', '#6366f1', '', ''],
        ]
        excel_file = create_excel_upload_file(headers, key_row, rows)
        res = self.client.post('/api/accounts/roles/import_excel/', {'file': excel_file, 'update_existing': 'true'}, format='multipart')
        self.assertEqual(res.status_code, 200)

        child_role.refresh_from_db()
        self.assertIsNone(child_role.parent)
        self.assertEqual(child_role.permissions.count(), 0)

    def test_atomic_rollback_on_role_import_error(self):
        headers = ['نام یکتا (name)', 'عنوان فارسی (title)', 'رنگ سازمانی (color)', 'نقش والد (parent)', 'مجوزها (permissions)']
        key_row = ['name', 'title', 'color', 'parent', 'permissions']
        rows = [
            ['atomic_role_1', 'نقش اتمیک ۱', '#4f46e5', '', ''],
            ['atomic_role_2', 'نقش اتمیک ۲', '#06b6d4', '', ''],
        ]
        excel_file = create_excel_upload_file(headers, key_row, rows)

        from unittest.mock import patch
        with patch.object(CustomRole, 'save', side_effect=[None, Exception("Fatal DB Error")]):
            res = self.client.post('/api/accounts/roles/import_excel/', {'file': excel_file}, format='multipart')
            self.assertEqual(res.status_code, 400)
            self.assertFalse(res.data['success'])
            self.assertIn('تراکنش کاملاً رول‌بک شد', res.data['errors'][0]['message'])

        self.assertFalse(CustomRole.objects.filter(name='atomic_role_1').exists())
        self.assertFalse(CustomRole.objects.filter(name='atomic_role_2').exists())


class UserAndRolePaginationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = CustomUser.objects.create_superuser('admin_pagination', 'admin@pagination.com', 'pass1234')
        self.client.force_authenticate(user=self.admin)
        for i in range(10):
            CustomUser.objects.create(
                username=f'pagination_user_{i}',
                first_name=f'کاربر{i}',
                last_name=f'تستی{i}',
                phone_number=f'0912000000{i}'
            )

    def test_users_unpaginated_by_default(self):
        # By default without ?page=, must return a list directly for dropdowns & client caches
        res = self.client.get('/api/accounts/users/')
        self.assertEqual(res.status_code, 200)
        self.assertIsInstance(res.data, list)
        self.assertGreaterEqual(len(res.data), 11)

    def test_users_paginated_when_page_param_provided(self):
        # When ?page= is provided, must return a paginated response with count & results
        res = self.client.get('/api/accounts/users/?page=1&page_size=5')
        self.assertEqual(res.status_code, 200)
        self.assertIsInstance(res.data, dict)
        self.assertIn('count', res.data)
        self.assertIn('results', res.data)
        self.assertEqual(len(res.data['results']), 5)
        self.assertGreaterEqual(res.data['count'], 11)

    def test_users_server_side_search(self):
        # Server-side search on username / first_name / phone
        res = self.client.get('/api/accounts/users/?search=pagination_user_3')
        self.assertEqual(res.status_code, 200)
        users = res.data if isinstance(res.data, list) else res.data['results']
        self.assertEqual(len(users), 1)
        self.assertEqual(users[0]['username'], 'pagination_user_3')

    def test_roles_paginated_when_page_param_provided(self):
        CustomRole.objects.create(name='test_role_p1', title='نقش تستی ۱')
        res = self.client.get('/api/accounts/roles/?page=1&page_size=10')
        self.assertEqual(res.status_code, 200)
        self.assertIsInstance(res.data, dict)
        self.assertIn('count', res.data)
        self.assertIn('results', res.data)


from accounts.models import UserDeviceSession
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.token_blacklist.models import BlacklistedToken, OutstandingToken

class UserSessionRevocationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = CustomUser.objects.create_superuser('admin_revoke_test', 'admin@revoke.com', 'pass1234')
        self.client.force_authenticate(user=self.admin)
        self.target_user = CustomUser.objects.create(
            username='user_to_reset',
            first_name='علی',
            last_name='تقوی',
            phone_number='09121234567'
        )
        self.session = UserDeviceSession.objects.create(
            user=self.target_user,
            session_key=f'{self.target_user.id}_tab1',
            tab_id='tab1',
            is_revoked=False
        )

    def test_admin_reset_password_revokes_sessions_and_blacklists_tokens(self):
        refresh = RefreshToken.for_user(self.target_user)
        self.assertFalse(self.session.is_revoked)

        res = self.client.post(f'/api/accounts/users/{self.target_user.id}/admin_reset_password/')
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data['success'])

        self.session.refresh_from_db()
        self.assertTrue(self.session.is_revoked)

        # Token must be blacklisted
        token_entry = OutstandingToken.objects.filter(token=str(refresh)).first()
        if token_entry:
            self.assertTrue(BlacklistedToken.objects.filter(token=token_entry).exists())

    def test_toggle_status_deactivation_revokes_sessions(self):
        self.assertFalse(self.session.is_revoked)
        res = self.client.patch(f'/api/accounts/users/{self.target_user.id}/toggle_status/', {'is_active': False})
        self.assertEqual(res.status_code, 200)

        self.session.refresh_from_db()
        self.assertTrue(self.session.is_revoked)


class ExcelFormulaInjectionSanitizationTests(TestCase):
    def test_sanitize_excel_cell_unit(self):
        from common.excel_utils import sanitize_excel_cell

        # Payload tests starting with dangerous characters
        self.assertEqual(sanitize_excel_cell("=cmd|' /C calc'!A0"), "'=cmd|' /C calc'!A0")
        self.assertEqual(sanitize_excel_cell("+1+1"), "'+1+1")
        self.assertEqual(sanitize_excel_cell("-10%"), "'-10%")
        self.assertEqual(sanitize_excel_cell("@SUM(1,2)"), "'@SUM(1,2)")
        self.assertEqual(sanitize_excel_cell("\t=calc"), "'\t=calc")
        self.assertEqual(sanitize_excel_cell("  =cmd"), "'  =cmd")

        # Normal text and numbers should not be modified
        self.assertEqual(sanitize_excel_cell("علی رضایی"), "علی رضایی")
        self.assertEqual(sanitize_excel_cell("پتروشیمی خلیج فارس"), "پتروشیمی خلیج فارس")
        self.assertEqual(sanitize_excel_cell(12345), 12345)
        self.assertEqual(sanitize_excel_cell(-99), -99)
        self.assertEqual(sanitize_excel_cell(3.14), 3.14)
        self.assertIsNone(sanitize_excel_cell(None))

        # Already sanitized strings should not be double escaped
        self.assertEqual(sanitize_excel_cell("'+123"), "'+123")

    def test_generate_users_excel_sanitizes_injected_formulas(self):
        import io
        from openpyxl import load_workbook
        from accounts.excel_utils import generate_users_excel

        user = CustomUser.objects.create(
            username="attacker_user",
            first_name="=cmd|' /C calc'!A0",
            last_name="+1+1",
            company="@EVIL_CORP",
            operational_zone="-ZONE-1",
            phone_number="09129998877"
        )

        response = generate_users_excel(CustomUser.objects.filter(id=user.id))
        self.assertEqual(response.status_code, 200)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        # Row 1 is Persian labels, Row 2 is keys, Row 3 is user data
        row3_values = [cell.value for cell in ws[3]]

        # Verify values in row 3 start with '
        first_name_val = row3_values[0]
        last_name_val = row3_values[1]
        zone_val = row3_values[8]
        company_val = row3_values[9]

        self.assertTrue(str(first_name_val).startswith("'="))
        self.assertTrue(str(last_name_val).startswith("'+"))
        self.assertTrue(str(zone_val).startswith("'-"))
        self.assertTrue(str(company_val).startswith("'@"))

    def test_generate_roles_excel_sanitizes_injected_formulas(self):
        import io
        from openpyxl import load_workbook
        from accounts.roles_excel_utils import generate_roles_excel
        from accounts.models import CustomRole
        role = CustomRole.objects.create(
            name="evil_role",
            title="=2+5+cmd|' /C calc'!A0",
            color="#ef4444"
        )

        response = generate_roles_excel(CustomRole.objects.filter(id=role.id))
        self.assertEqual(response.status_code, 200)

        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        row3_values = [cell.value for cell in ws[3]]
        title_val = row3_values[1]
        self.assertTrue(str(title_val).startswith("'="))


