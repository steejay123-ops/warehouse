from django.test import TestCase
from django.contrib.auth import get_user_model
from decimal import Decimal
from rest_framework.test import APIClient
from personnel.models import (
    FinancialProject,
    ProjectSection,
    UserSectionAssignment,
    Counterparty,
    ExpenseInvoice,
    PersonnelProfile,
    VehicleDriverProfile,
    DailyAttendance,
    VehicleTripLog
)
from personnel.section_guardian import SectionGuardian

User = get_user_model()


class SectionGuardianPhase1Tests(TestCase):
    """
    مجموعه تست‌های خودکار ایجنت نگهبان فونداسیون معماری پروژه و بخش
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username='test_operator',
            password='Password123!',
            first_name='علی',
            last_name='رضایی'
        )
        self.supervisor = User.objects.create_user(
            username='test_supervisor',
            password='Password123!',
            first_name='مهندس',
            last_name='احمدی'
        )
        self.project = FinancialProject.objects.create(
            code='PRJ-SHIRAZ',
            name='پروژه انبار مرکزی شیراز',
            description='انبار لجستیک جنوب کشور'
        )
        self.section = ProjectSection.objects.create(
            project=self.project,
            code='SEC-TRANS',
            name='بخش ترابری و لجستیک'
        )

    def test_section_guardian_automated_audit(self):
        """اجرای ممیزی خودکار ایجنت نگهبان"""
        guardian = SectionGuardian()
        self.assertTrue(guardian.audit_phase_1_foundation())

    def test_financial_project_and_section_creation(self):
        """تست ایجاد پروژه و بخش و روابط یکتایی"""
        self.assertEqual(str(self.project), "پروژه انبار مرکزی شیراز (PRJ-SHIRAZ)")
        self.assertEqual(str(self.section), "پروژه انبار مرکزی شیراز - بخش ترابری و لجستیک (SEC-TRANS)")
        self.assertEqual(self.project.sections.count(), 1)
        self.assertEqual(self.project.sections.first(), self.section)

    def test_user_section_assignment(self):
        """تست انتساب کاربر به بخش با نقش‌های مجاز"""
        assignment = UserSectionAssignment.objects.create(
            user=self.user,
            section=self.section,
            role='employee'
        )
        self.assertEqual(assignment.role, 'employee')
        self.assertTrue(assignment.is_active)
        self.assertEqual(self.user.section_assignments.count(), 1)

    def test_counterparty_creation(self):
        """تست ایجاد طرف‌حساب به عنوان موجودیت مالی سراسری مستقل"""
        cp = Counterparty.objects.create(
            name='تعمیرگاه تخصصی دیزل پارس',
            counterparty_type='repair_shop',
            phone='09171234567'
        )
        self.assertEqual(cp.counterparty_type, 'repair_shop')
        self.assertIn("تعمیرگاه", str(cp))

    def test_expense_invoice_draft_invariant(self):
        """تست الزام وضعیت پیش‌نویس (draft) برای فاکتور جدید"""
        cp = Counterparty.objects.create(
            name='جایگاه سوخت پارسیان',
            counterparty_type='fuel_station'
        )
        invoice = ExpenseInvoice.objects.create(
            section=self.section,
            counterparty=cp,
            invoice_number='INV-1404-001',
            invoice_date_shamsi='1404/05/10',
            amount=Decimal('5500000'),
            category='سوخت ناوگان',
            description='بنزین و گازوئیل خودروهای شیفت روز',
            created_by=self.user
        )
        self.assertEqual(invoice.status, 'draft', "وضعیت اولیه فاکتور باید اجباراً draft باشد")
        self.assertEqual(invoice.amount, Decimal('5500000'))
        self.assertIn("INV-1404-001", str(invoice))

    def test_backward_compatibility_existing_models_without_project(self):
        """تست عدم وابستگی اجباری و کارکرد کامل رکوردهای قدیمی بدون پروژه و بخش"""
        personnel = PersonnelProfile.objects.create(
            first_name='حسن',
            last_name='مرادی',
            national_code='1234567890',
            job_title='کارشناس انبار'
        )
        self.assertIsNone(personnel.project)
        self.assertIsNone(personnel.section)

        # انتساب اختیاری
        personnel.project = self.project
        personnel.section = self.section
        personnel.save()
        self.assertEqual(personnel.project, self.project)
        self.assertEqual(personnel.section, self.section)

        # کارکرد روزانه پرسنل
        attendance = DailyAttendance.objects.create(
            personnel=personnel,
            date_shamsi='1404/05/12',
            project=self.project,
            section=self.section,
            effective_hours=Decimal('10.00')
        )
        self.assertEqual(attendance.section, self.section)


class SectionGuardianPhase2APITests(TestCase):
    """
    تست‌های یکپارچگی اندپوینت‌های API فاز ۲ (مدیریت پروژه، بخش و انتساب کاربران)
    """

    def setUp(self):
        self.client = APIClient()
        self.admin_user = User.objects.create_superuser(
            username='admin_boss',
            password='AdminPassword123!',
            email='admin@company.com'
        )
        self.employee_user = User.objects.create_user(
            username='employee_ali',
            password='Password123!'
        )
        self.client.force_authenticate(user=self.admin_user)

    def test_api_crud_financial_project(self):
        """تست ایجاد و دریافت پروژه از طریق API"""
        res = self.client.post('/api/personnel/financial-projects/', {
            'code': 'PRJ-TEHRAN',
            'name': 'پروژه پایانه تهران غرب',
            'description': 'انبار توزیع مکانیزه پایتخت'
        }, format='json')
        self.assertEqual(res.status_code, 201)
        project_id = res.data['id']

        # لیست پروژه‌ها
        get_res = self.client.get('/api/personnel/financial-projects/')
        self.assertEqual(get_res.status_code, 200)
        self.assertTrue(any(p['code'] == 'PRJ-TEHRAN' for p in get_res.data))

        # ایجاد بخش ذیل پروژه
        sec_res = self.client.post('/api/personnel/project-sections/', {
            'project': project_id,
            'code': 'SEC-WAREHOUSE',
            'name': 'بخش تفکیک و بسته‌بندی'
        }, format='json')
        self.assertEqual(sec_res.status_code, 201)
        sec_id = sec_res.data['id']

        # انتساب کاربر به بخش
        assign_res = self.client.post('/api/personnel/user-section-assignments/', {
            'user': self.employee_user.id,
            'section': sec_id,
            'role': 'employee'
        }, format='json')
        self.assertEqual(assign_res.status_code, 201)
        self.assertEqual(assign_res.data['role'], 'employee')

        # لاگین به عنوان کارمند و دریافت my-sections
        self.client.force_authenticate(user=self.employee_user)
        my_sec_res = self.client.get('/api/personnel/user-section-assignments/my-sections/')
        self.assertEqual(my_sec_res.status_code, 200)
        self.assertEqual(len(my_sec_res.data), 1)
        self.assertEqual(my_sec_res.data[0]['code'], 'SEC-WAREHOUSE')

    def test_api_counterparty_and_invoice(self):
        """تست تعریف طرف‌حساب و صدور فاکتور هزینه با قفل draft"""
        proj = FinancialProject.objects.create(code='PRJ-ISFAHAN', name='پروژه اصفهان')
        sec = ProjectSection.objects.create(project=proj, code='SEC-FLEET', name='ناوگان اصفهان')

        # تعریف طرف‌حساب
        cp_res = self.client.post('/api/personnel/counterparties/', {
            'name': 'فروشگاه لوازم یدکی سپاهان',
            'counterparty_type': 'repair_shop',
            'phone': '03131234567'
        }, format='json')
        self.assertEqual(cp_res.status_code, 201)
        cp_id = cp_res.data['id']

        # ثبت فاکتور
        inv_res = self.client.post('/api/personnel/expense-invoices/', {
            'section': sec.id,
            'counterparty': cp_id,
            'invoice_number': 'FAK-9901',
            'invoice_date_shamsi': '1404/05/15',
            'amount': 12500000,
            'category': 'تعمیرات و تعویض روغن',
            'description': 'سرویس دوره‌ای کامیون‌ها'
        }, format='json')
        self.assertEqual(inv_res.status_code, 201)
        self.assertEqual(inv_res.data['status'], 'draft', "فاکتور باید در وضعیت draft ایجاد شود")

    def test_clone_assignments_api(self):
        """تست اختصاصی تکثیر و کپی چارت سازمانی از بخش مبدأ به مقصد"""
        proj = FinancialProject.objects.create(code='PRJ-CLONE', name='پروژه تکثیر چارت')
        sec_source = ProjectSection.objects.create(project=proj, code='SEC-SRC', name='بخش الگو')
        sec_target = ProjectSection.objects.create(project=proj, code='SEC-TGT', name='بخش مقصد')

        # انتساب کاربر به بخش مبدأ
        UserSectionAssignment.objects.create(
            user=self.admin_user,
            section=sec_source,
            role='manager',
            is_active=True
        )
        UserSectionAssignment.objects.create(
            user=self.employee_user,
            section=sec_source,
            role='employee',
            is_active=True
        )

        # فراخوانی اکشن کپی چارت
        url = f'/api/personnel/project-sections/{sec_target.id}/clone-assignments/'
        res = self.client.post(url, {'source_section_id': sec_source.id}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.data['success'])
        self.assertEqual(res.data['copied_count'], 2)

        # بررسی وجود انتساب‌ها در بخش مقصد
        target_assignments = UserSectionAssignment.objects.filter(section=sec_target)
        self.assertEqual(target_assignments.count(), 2)
        roles = set(target_assignments.values_list('role', flat=True))
        self.assertEqual(roles, {'manager', 'employee'})

        # فراخوانی مجدد (باید انتساب تکراری اضافه نکند)
        res_duplicate = self.client.post(url, {'source_section_id': sec_source.id}, format='json')
        self.assertEqual(res_duplicate.status_code, 200)
        self.assertEqual(res_duplicate.data['copied_count'], 0)
        self.assertEqual(res_duplicate.data['skipped_count'], 2)

    def test_excel_export_and_template_endpoints(self):
        """تست اندپوینت‌های صدور و قالب اکسل برای پروژه‌ها، بخش‌ها و انتساب‌ها"""
        # پروژه‌ها
        res_prj_exp = self.client.get('/api/personnel/financial-projects/export-excel/')
        self.assertEqual(res_prj_exp.status_code, 200)
        self.assertIn('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', res_prj_exp['Content-Type'])

        res_prj_tpl = self.client.get('/api/personnel/financial-projects/download-template/')
        self.assertEqual(res_prj_tpl.status_code, 200)

        # بخش‌ها
        res_sec_exp = self.client.get('/api/personnel/project-sections/export-excel/')
        self.assertEqual(res_sec_exp.status_code, 200)

        res_sec_tpl = self.client.get('/api/personnel/project-sections/download-template/')
        self.assertEqual(res_sec_tpl.status_code, 200)

        # انتساب‌ها
        res_asn_exp = self.client.get('/api/personnel/user-section-assignments/export-excel/')
        self.assertEqual(res_asn_exp.status_code, 200)

        res_asn_tpl = self.client.get('/api/personnel/user-section-assignments/download-template/')
        self.assertEqual(res_asn_tpl.status_code, 200)

    def test_bulk_assign_api(self):
        """تست اختصاصی انتساب دسته‌جمعی چند کاربر به یک بخش و نقش"""
        proj = FinancialProject.objects.create(code='PRJ-BULK', name='پروژه تست گروهی')
        sec = ProjectSection.objects.create(project=proj, code='SEC-BLK', name='بخش انتساب چندگانه')

        # درخواست انتساب همزمان مدیر و کارمند
        url = '/api/personnel/user-section-assignments/bulk-assign/'
        payload = {
            'section_id': sec.id,
            'user_ids': [self.admin_user.id, self.employee_user.id],
            'role': 'accountant'
        }
        res = self.client.post(url, payload, format='json')
        self.assertEqual(res.status_code, 201)
        self.assertEqual(len(res.data), 2)
        self.assertTrue(all(a['role'] == 'accountant' for a in res.data))
        self.assertTrue(all(a['section'] == sec.id for a in res.data))

        # بررسی در دیتابیس
        count = UserSectionAssignment.objects.filter(section=sec, role='accountant', is_active=True).count()
        self.assertEqual(count, 2)

    def test_soft_deactivation_on_destroy_with_invoices(self):
        """تست عدم حذف فیزیکی و تبدیل به غیرفعال‌سازی نرم در صورت وجود فاکتور یا سند مالی"""
        proj = FinancialProject.objects.create(code='PRJ-SAFE', name='پروژه حفاظت داده')
        sec = ProjectSection.objects.create(project=proj, code='SEC-SAFE', name='بخش امن')
        cp = Counterparty.objects.create(name='تأمین‌کننده قطعات امین')
        ExpenseInvoice.objects.create(
            section=sec,
            counterparty=cp,
            invoice_number='INV-SAFE-01',
            invoice_date_shamsi='1404/06/01',
            amount=1000000,
            category='مصرفی',
            description='خرید قطعات ضروری'
        )

        # تلاش برای حذف بخش از طریق API
        del_sec_res = self.client.delete(f'/api/personnel/project-sections/{sec.id}/')
        self.assertEqual(del_sec_res.status_code, 204)
        sec.refresh_from_db()
        self.assertFalse(sec.is_active, "بخش دارای فاکتور نباید فیزیکی حذف شود بلکه باید is_active=False گردد")

        # تلاش برای حذف طرف‌حساب از طریق API
        del_cp_res = self.client.delete(f'/api/personnel/counterparties/{cp.id}/')
        self.assertEqual(del_cp_res.status_code, 204)
        cp.refresh_from_db()
        self.assertFalse(cp.is_active, "طرف‌حساب دارای فاکتور نباید فیزیکی حذف شود بلکه باید is_active=False گردد")

        # تلاش برای حذف پروژه از طریق API
        del_proj_res = self.client.delete(f'/api/personnel/financial-projects/{proj.id}/')
        self.assertEqual(del_proj_res.status_code, 204)
        proj.refresh_from_db()
        self.assertFalse(proj.is_active, "پروژه دارای فاکتور نباید فیزیکی حذف شود بلکه باید is_active=False گردد")

    def test_reactivate_inactive_assignment_without_unique_error(self):
        """تست فعال‌سازی مجدد رکورد غیرفعال انتساب بدون خطای یکتایی دیتابیس"""
        proj = FinancialProject.objects.create(code='PRJ-REACT', name='پروژه فعال‌سازی مجدد')
        sec = ProjectSection.objects.create(project=proj, code='SEC-REACT', name='بخش فعال‌سازی')
        assignment = UserSectionAssignment.objects.create(
            user=self.employee_user,
            section=sec,
            role='supervisor',
            is_active=False
        )

        # ارسال مجدد ایجاد انتساب از طریق API
        res = self.client.post('/api/personnel/user-section-assignments/', {
            'user': self.employee_user.id,
            'section': sec.id,
            'role': 'supervisor'
        }, format='json')
        self.assertEqual(res.status_code, 201)
        assignment.refresh_from_db()
        self.assertTrue(assignment.is_active, "انتساب غیرفعال قبلی باید بدون ارور مجدداً فعال شود")

    def test_dry_run_excel_import(self):
        """تست عملکرد حالت dry_run در متدهای اکسل بدون ذخیره در دیتابیس"""
        import openpyxl
        import io
        from personnel.org_excel_engine import import_projects_from_excel

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="کد پروژه")
        ws.cell(row=2, column=1, value="code")
        ws.cell(row=3, column=1, value="PRJ-DRY99")
        ws.cell(row=3, column=2, value="پروژه آزمایشی درای‌ران")
        ws.cell(row=3, column=3, value="توضیح")
        ws.cell(row=3, column=4, value="بله")

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        # اجرای dry_run
        res = import_projects_from_excel(buf, dry_run=True)
        self.assertTrue(res['success'])
        self.assertTrue(res['dry_run'])
        self.assertEqual(res['summary']['created'], 1)
        # اطمینان از عدم ایجاد در دیتابیس
        self.assertFalse(FinancialProject.objects.filter(code='PRJ-DRY99').exists())

    def test_counterparty_sheba_validation(self):
        """تست اعتبارسنجی سریالایزر طرف‌حساب برای شبا و کد ملی"""
        from personnel.serializers import CounterpartySerializer

        # شبای نامعتبر
        ser_invalid = CounterpartySerializer(data={
            'name': 'طرف‌حساب تست شبا',
            'sheba_number': '123456789012345678901234' # چک‌سام اشتباه
        })
        self.assertFalse(ser_invalid.is_valid())
        self.assertIn('sheba_number', ser_invalid.errors)

        # شبای معتبر بانک ملی
        ser_valid = CounterpartySerializer(data={
            'name': 'طرف‌حساب تست شبا معتبر',
            'sheba_number': '780170000000101111111001',
            'national_id': '1234567890'
        })
        self.assertTrue(ser_valid.is_valid(), ser_valid.errors)



