from django.test import TestCase
from rest_framework.test import APIClient
from django.contrib.auth.models import Permission
from decimal import Decimal
from accounts.models import CustomUser
from personnel.models import (
    FinancialProject,
    ProjectSection,
    VehicleDriverProfile,
    VehicleChangeRequest,
    VehicleTripLog,
    UserSectionAssignment,
)
from personnel.sheba_utils import generate_sheba_from_account


class VehicleLifecycleTests(TestCase):
    """
    تست‌های جامع چرخه حیات ناوگان، امنیت قلمرو، جلوگیری از تنزل رتبه در تغییرات و حذف امن
    """

    def setUp(self):
        self.project = FinancialProject.objects.create(code="1", name="پروژه دالان", is_active=True)
        self.section_a = ProjectSection.objects.create(project=self.project, code="11", name="کارگاه دالان", is_active=True)
        self.section_b = ProjectSection.objects.create(project=self.project, code="12", name="انبار دالان", is_active=True)

        self.operator = CustomUser.objects.create(username="veh_operator", first_name="رضا", last_name="اکبری")
        self.supervisor = CustomUser.objects.create(username="veh_supervisor", first_name="حمید", last_name="صالحی")
        self.accountant = CustomUser.objects.create(username="veh_accountant", first_name="مهری", last_name="باقری")
        self.manager = CustomUser.objects.create(username="veh_manager", first_name="سید", last_name="موسوی")
        self.outsider = CustomUser.objects.create(username="veh_outsider", first_name="کاربر", last_name="بی‌دسترسی")

        # انتساب مجوزها
        for codename in ["perm_approve_fleet_supervisor", "can_act_as_workshop_supervisor"]:
            p = Permission.objects.filter(codename=codename).first()
            if p:
                self.supervisor.user_permissions.add(p)

        p_fin = Permission.objects.filter(codename="perm_approve_fleet_finance").first()
        if p_fin:
            self.accountant.user_permissions.add(p_fin)

        for codename in ["perm_approve_fleet_manager", "can_act_as_manager"]:
            p = Permission.objects.filter(codename=codename).first()
            if p:
                self.manager.user_permissions.add(p)

        # انتساب بخش A به کاربران مجاز
        UserSectionAssignment.objects.create(user=self.operator, section=self.section_a, role="employee", is_active=True)
        UserSectionAssignment.objects.create(user=self.supervisor, section=self.section_a, role="supervisor", is_active=True)
        UserSectionAssignment.objects.create(user=self.accountant, section=self.section_a, role="accountant", is_active=True)
        UserSectionAssignment.objects.create(user=self.manager, section=self.section_a, role="manager", is_active=True)

        self.client = APIClient()

    def test_complete_vehicle_approval_lifecycle(self):
        """
        تست چرخه کامل: ثبت پیش‌نویس -> تایید سرپرست -> تایید مالی -> تصویب نهایی مدیر
        """
        valid_sheba = generate_sheba_from_account("بانک ملی ایران", "0105678901004")
        self.client.force_authenticate(user=self.operator)

        # ۱. ثبت با وضعیت پیش‌نویس
        payload = {
            "plate_number": "12الف345ایران63",
            "vehicle_type": "nissan",
            "ownership_type": "contract",
            "driver_name": "احمد قاسمی",
            "driver_national_code": "0010376488",
            "driver_phone": "09121234567",
            "default_service_rate": 4500000,
            "bank_name": "بانک ملی ایران",
            "sheba_number": valid_sheba,
            "section": self.section_a.id,
            "approval_status": "draft",
        }
        res_create = self.client.post("/api/personnel/vehicles/", payload, format="json")
        self.assertEqual(res_create.status_code, 201)
        veh_id = res_create.data["id"]
        self.assertEqual(res_create.data["approval_status"], "draft")

        # ۲. اپراتور ارسال به سرپرست می‌کند
        res_submit = self.client.patch(f"/api/personnel/vehicles/{veh_id}/", {"approval_status": "pending_supervisor"}, format="json")
        self.assertEqual(res_submit.status_code, 200)
        self.assertEqual(res_submit.data["approval_status"], "pending_supervisor")

        # ۳. تایید سرپرست
        self.client.force_authenticate(user=self.supervisor)
        res_sup = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-supervisor/")
        self.assertEqual(res_sup.status_code, 200)
        self.assertEqual(res_sup.data["data"]["approval_status"], "pending_accountant")

        # ۴. تایید مالی
        self.client.force_authenticate(user=self.accountant)
        res_fin = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-finance/")
        self.assertEqual(res_fin.status_code, 200)
        self.assertEqual(res_fin.data["data"]["approval_status"], "pending_manager")

        # ۵. تصویب نهایی مدیر
        self.client.force_authenticate(user=self.manager)
        res_mgr = self.client.post(f"/api/personnel/vehicles/{veh_id}/approve-manager/")
        self.assertEqual(res_mgr.status_code, 200)
        self.assertEqual(res_mgr.data["data"]["approval_status"], "approved")
        self.assertTrue(res_mgr.data["data"]["is_active"])

    def test_prevent_demotion_bug_on_change_request_approval(self):
        """
        راستی‌آزمایی رفع باگ تنزل رتبه:
        تایید درخواست تغییرات یک خودروی مصوب، نباید وضعیت خودرو را به pending یا draft بازگرداند!
        """
        # ایجاد یک خودروی تاییدشده اولیه
        vehicle = VehicleDriverProfile.objects.create(
            plate_number="34ب567ایران63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="محسن رضایی",
            driver_national_code="0010376488",
            default_service_rate=Decimal("6000000"),
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )

        # اپراتور درخواست تغییر نرخ و نام راننده ارسال می‌کند
        self.client.force_authenticate(user=self.operator)
        edit_payload = {
            "driver_name": "محسن رضایی اصل",
            "default_service_rate": 6500000,
            "approval_status": "pending_supervisor",  # حتی اگر فرانت‌اند وضعیت را بفرستد
        }
        res_edit = self.client.patch(f"/api/personnel/vehicles/{vehicle.id}/", edit_payload, format="json")
        self.assertEqual(res_edit.status_code, 202)  # درخواست به کارتابل تغییرات رفت
        cr_id = res_edit.data["change_request_id"]

        cr = VehicleChangeRequest.objects.get(id=cr_id)
        self.assertEqual(cr.status, "pending_supervisor")

        # مدیر مستقیماً درخواست تغییرات را تایید می‌کند
        self.client.force_authenticate(user=self.manager)
        res_approve_cr = self.client.post(f"/api/personnel/vehicle-change-requests/{cr_id}/approve-manager/")
        self.assertEqual(res_approve_cr.status_code, 200)

        # بررسی پرونده اصلی خودرو در دیتابیس
        vehicle.refresh_from_db()
        self.assertEqual(vehicle.driver_name, "محسن رضایی اصل")
        self.assertEqual(vehicle.default_service_rate, Decimal("6500000"))
        # مهم‌ترین اعتبارسنجی: وضعیت خودرو همچنان باید مصوب باقی بماند و به pending تنزل نیابد!
        self.assertEqual(vehicle.approval_status, "approved")
        self.assertFalse(vehicle.has_pending_changes)

    def test_section_scoping_and_bola_protection(self):
        """
        بررسی ایزولاسیون قلمرو (Guardian G1): کاربر بدون انتساب نباید به خودروهای بخش دیگر دسترسی داشته باشد
        """
        VehicleDriverProfile.objects.create(
            plate_number="99ج999ایران99",
            vehicle_type="nissan",
            driver_name="راننده بخش B",
            section=self.section_b,
            approval_status="approved",
            is_active=True
        )

        # کاربر outsider دسترسی به بخش B ندارد
        self.client.force_authenticate(user=self.outsider)
        res_list = self.client.get(f"/api/personnel/vehicles/?section_id={self.section_b.id}")
        self.assertEqual(len(res_list.data), 0)

        # کاربر outsider نمی‌تواند در بخش B خودرو ثبت کند
        create_payload = {
            "plate_number": "88د888ایران88",
            "vehicle_type": "nissan",
            "driver_name": "تلاش غیرمجاز",
            "section": self.section_b.id,
        }
        res_create = self.client.post("/api/personnel/vehicles/", create_payload, format="json")
        self.assertEqual(res_create.status_code, 403)

    def test_safe_destroy_prevents_hard_cascade_delete(self):
        """
        بررسی حذف امن: خودرویی که دارای سابقه کارکرد و تردد است نباید حذف فیزیکی شود، بلکه نرم غیرفعال می‌گردد
        """
        vehicle = VehicleDriverProfile.objects.create(
            plate_number="55س555ایران55",
            vehicle_type="nissan",
            driver_name="راننده دارای سابقه",
            section=self.section_a,
            approval_status="draft",
            is_active=True
        )
        # ثبت یک تردد برای این خودرو
        VehicleTripLog.objects.create(
            vehicle=vehicle,
            date_shamsi="1403/04/01",
            trip_count=2,
            unit_rate=4000000,
            total_amount=8000000
        )

        self.client.force_authenticate(user=self.operator)
        res_del = self.client.delete(f"/api/personnel/vehicles/{vehicle.id}/")
        self.assertEqual(res_del.status_code, 200)

        # رکورد نباید از دیتابیس پاک شده باشد، فقط is_active = False
        vehicle.refresh_from_db()
        self.assertFalse(vehicle.is_active)
        self.assertTrue(VehicleDriverProfile.objects.filter(id=vehicle.id).exists())

    def test_excel_download_template(self):
        """
        دانلود قالب اکسل ناوگان
        """
        self.client.force_authenticate(user=self.operator)
        res = self.client.get("/api/personnel/vehicles/download-template/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_excel_import_lifecycle_vehicles(self):
        """
        تست چرخه کامل درون‌ریزی اکسل ناوگان: دانلود قالب، ارسال در حالت پیش‌نمایش و اعمال قطعی
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_authenticate(user=self.operator)

        # ۱. دریافت قالب اکسل
        tpl_res = self.client.get("/api/personnel/vehicles/download-template/")
        self.assertEqual(tpl_res.status_code, 200)

        # ۲. پیش‌نمایش (Dry-Run)
        excel_file = SimpleUploadedFile("template.xlsx", tpl_res.content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res_dry = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file, "section_id": self.section_a.id, "dry_run": "true"},
            format="multipart"
        )
        self.assertEqual(res_dry.status_code, 200)
        self.assertTrue(res_dry.data["success"])
        self.assertTrue(res_dry.data["dry_run"])
        self.assertEqual(res_dry.data["summary"]["created"], 1)
        self.assertFalse(VehicleDriverProfile.objects.filter(plate_number="12 الف 345 ایران 63").exists())

        # ۳. اعمال قطعی
        excel_file_real = SimpleUploadedFile("template.xlsx", tpl_res.content, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res_commit = self.client.post(
            "/api/personnel/vehicles/import-excel/",
            {"file": excel_file_real, "section_id": self.section_a.id, "dry_run": "false"},
            format="multipart"
        )
        self.assertEqual(res_commit.status_code, 200)
        self.assertTrue(res_commit.data["success"])
        self.assertFalse(res_commit.data["dry_run"])
        self.assertEqual(res_commit.data["created_count"], 1)

        # بررسی در دیتابیس
        v = VehicleDriverProfile.objects.filter(plate_number="12 الف 345 ایران 63").first()
        self.assertIsNotNone(v)
        self.assertEqual(v.driver_name, "رضا اکبری")
        self.assertEqual(v.approval_status, "draft")
        self.assertEqual(v.section_id, self.section_a.id)

    def test_vehicle_owner_fields_and_validation(self):
        """
        تست تفکیک راننده و مالک خودرو و اعتبارسنجی کدملی مالک
        """
        self.client.force_authenticate(user=self.operator)

        # ۱. تست ثبت موفق با مالک مجزا
        payload = {
            "plate_number": "99ط888ایران63",
            "vehicle_type": "truck",
            "ownership_type": "contract",
            "driver_name": "سعید رضایی",
            "driver_national_code": "0010376488",
            "is_driver_owner": False,
            "owner_name": "حسین حسینی",
            "owner_national_code": "0010376488",
            "owner_phone": "09123334455",
            "section": self.section_a.id,
            "approval_status": "draft"
        }
        res = self.client.post("/api/personnel/vehicles/", payload, format="json")
        self.assertEqual(res.status_code, 201)
        v = VehicleDriverProfile.objects.get(id=res.data["id"])
        self.assertFalse(v.is_driver_owner)
        self.assertEqual(v.owner_name, "حسین حسینی")
        self.assertEqual(v.owner_national_code, "0010376488")
        self.assertEqual(v.owner_phone, "09123334455")

        # ۲. تست رد شدن کدملی نامعتبر مالک
        bad_payload = {
            "plate_number": "88ق777ایران63",
            "vehicle_type": "nissan",
            "ownership_type": "contract",
            "driver_name": "سعید رضایی",
            "is_driver_owner": False,
            "owner_name": "حسین حسینی",
            "owner_national_code": "1111111111",
            "section": self.section_a.id,
            "approval_status": "draft"
        }
        res_bad = self.client.post("/api/personnel/vehicles/", bad_payload, format="json")
        self.assertEqual(res_bad.status_code, 400)
        self.assertIn("owner_national_code", res_bad.data)

    def test_export_excel_includes_owner_columns(self):
        """
        تست وجود ستون‌های مالک و نوع مالکیت در خروجی رسمی اکسل ناوگان
        """
        import io
        import openpyxl

        self.client.force_authenticate(user=self.operator)
        VehicleDriverProfile.objects.create(
            plate_number="55 ع 333 ایران 63",
            vehicle_type="khavar",
            ownership_type="contract",
            driver_name="محمود کاظمی",
            driver_national_code="0010376488",
            driver_phone="09121112233",
            is_driver_owner=False,
            owner_name="جواد احمدی",
            owner_national_code="0078901235",
            owner_phone="09124445566",
            default_service_rate=Decimal("2500000"),
            section=self.section_a,
            approval_status="approved",
            is_active=True
        )

        res = self.client.get(f"/api/personnel/vehicles/export-excel/?section_id={self.section_a.id}")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "ناوگان و رانندگان")

        # بررسی سطرهای هدر
        headers_fa = [cell.value for cell in ws[1]]
        headers_en = [cell.value for cell in ws[2]]

        self.assertIn("مالک شخص راننده است", headers_fa)
        self.assertIn("نام مالک", headers_fa)
        self.assertIn("کد ملی مالک", headers_fa)
        self.assertIn("شماره همراه مالک", headers_fa)
        self.assertIn("نوع مالکیت", headers_fa)

        self.assertIn("is_driver_owner", headers_en)
        self.assertIn("owner_name", headers_en)
        self.assertIn("owner_national_code", headers_en)
        self.assertIn("owner_phone", headers_en)
        self.assertIn("ownership_type_display", headers_en)

        # بررسی مقادیر ردیف داده
        data_row = [cell.value for cell in ws[3]]
        self.assertIn("55 ع 333 ایران 63", data_row)
        self.assertIn("خیر", data_row)
        self.assertIn("جواد احمدی", data_row)
        self.assertIn("0078901235", data_row)
        self.assertIn("09124445566", data_row)
